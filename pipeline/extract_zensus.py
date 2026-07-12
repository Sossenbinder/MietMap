"""Extract Gemeinde- and Kreis-level housing metrics from the Zensus 2022 Regionaltabelle.

Input:  data/Regionaltabelle_Gebaeude_Wohnungen.xlsx (sheet CSV-Wohnungen)
Output: data/out/zensus_gemeinden.csv  (ars, name, level, qm_miete, leerstand, eigentum)
        data/out/zensus_kreise.csv     (kreis key = 5-digit ARS prefix, for fallback fills)
"""

from pathlib import Path

import openpyxl
import pandas as pd

DATA = Path(__file__).parent / "data"
OUT = DATA / "out"

COLS = {
    "_RS": "ars",
    "Name": "name",
    "Regionalebene": "level",
    "QMMIETE": "qm_miete",
    "LEQ": "leerstand",
    "ETQ": "eigentum",
}


def read_sheet() -> pd.DataFrame:
    wb = openpyxl.load_workbook(DATA / "Regionaltabelle_Gebaeude_Wohnungen.xlsx", read_only=True)
    ws = wb["CSV-Wohnungen"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {name: header.index(name) for name in COLS}
    records = [[r[idx[name]] for name in COLS] for r in rows]
    wb.close()
    return pd.DataFrame(records, columns=list(COLS.values()))


def to_num(s: pd.Series) -> pd.Series:
    # German decimal commas + privacy placeholders ("–", ".", "x") → NaN
    return pd.to_numeric(
        s.astype(str).str.replace(",", ".", regex=False), errors="coerce"
    )


def main() -> None:
    OUT.mkdir(exist_ok=True)
    df = read_sheet()
    for c in ("qm_miete", "leerstand", "eigentum"):
        df[c] = to_num(df[c])

    gem = df[df["level"] == "Gemeinde"].drop(columns="level").copy()
    gem["ars"] = gem["ars"].astype(str).str.zfill(12)

    kreis = df[df["level"] == "Stadtkreis/kreisfreie Stadt/Landkreis"].drop(columns="level").copy()
    kreis["ars"] = kreis["ars"].astype(str).str.zfill(5)

    gem.to_csv(OUT / "zensus_gemeinden.csv", index=False)
    kreis.to_csv(OUT / "zensus_kreise.csv", index=False)
    print(f"gemeinden: {len(gem)} (qm_miete missing: {gem.qm_miete.isna().sum()})")
    print(f"kreise:    {len(kreis)}")


if __name__ == "__main__":
    main()
